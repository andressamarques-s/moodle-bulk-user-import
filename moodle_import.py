"""Business rules for bulk user import into Moodle.

Takes a batch of people (organization, unit, name, CPF, e-mail, course), validates it and
generates three files:

* CSV  - ready for Moodle's "Upload users" tool
* PDF  - credentials list to hand out
* XLSX - workbook with the data, the PDF list and the CSV columns
"""
import csv
import io
import os
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

CSV_HEADER = ["username", "firstname", "lastname", "email", "password",
              "institution", "course1", "group1"]
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@dataclass
class Person:
    org: str
    unit: str
    name: str
    cpf: str
    email: str
    course: str
    errors: list = field(default_factory=list)

    @property
    def firstname(self):
        return self.name.split(" ", 1)[0]

    @property
    def lastname(self):
        parts = self.name.split(" ", 1)
        return parts[1] if len(parts) > 1 else ""


def clean(text):
    """Strip quotes, commas, semicolons and line breaks (they would break the CSV)."""
    if text is None:
        return ""
    text = str(text).replace('"', "").replace(",", " ").replace(";", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize_cpf(value):
    """Digits only, left-padded with zeros to 11 (like Excel's TEXT(x;"00000000000"))."""
    if value is None:
        return ""
    if isinstance(value, float):
        value = int(value)
    digits = re.sub(r"\D", "", str(value))
    return digits.zfill(11) if digits else ""


def is_valid_cpf(cpf):
    """Validates the two check digits of a Brazilian CPF."""
    if len(cpf) != 11 or cpf == cpf[0] * 11:
        return False
    for size in (9, 10):
        total = sum(int(cpf[i]) * (size + 1 - i) for i in range(size))
        digit = (total * 10 % 11) % 10
        if digit != int(cpf[size]):
            return False
    return True


def build_people(rows, valid_courses):
    """Normalize and validate the rows sent by the form. Fully empty rows are skipped."""
    people = []
    for row in rows:
        p = Person(
            org=clean(row.get("org")).upper(),
            unit=clean(row.get("unit")),
            name=clean(row.get("name")),
            cpf=normalize_cpf(row.get("cpf")),
            email=clean(row.get("email")).replace(" ", ""),
            course=clean(row.get("course")),
        )
        # the org code comes from the top of the form on every row, so it doesn't count as "filled"
        if not any([p.unit, p.name, p.cpf, p.email, p.course]):
            continue
        for attr, label in [("org", "Organization code"), ("unit", "Unit"),
                            ("name", "Full name"), ("cpf", "CPF"),
                            ("email", "E-mail"), ("course", "Course")]:
            if not getattr(p, attr):
                p.errors.append(f"Fill in: {label}")
        if p.name and " " not in p.name:
            p.errors.append("Enter first and last name")
        if p.cpf and not is_valid_cpf(p.cpf):
            p.errors.append("Invalid CPF")
        if p.email and not EMAIL_RE.match(p.email):
            p.errors.append("Invalid e-mail")
        if p.course and valid_courses and p.course not in valid_courses:
            p.errors.append("Course is not in the Moodle course list")
        people.append(p)

    # Batch rules: single organization, no duplicated CPF
    orgs = {p.org for p in people if p.org}
    if len(orgs) > 1:
        main = people[0].org
        for p in people:
            if p.org and p.org != main:
                p.errors.append(f"All rows must belong to the same organization ({main})")
    seen = {}
    for number, p in enumerate(people, start=1):
        if p.cpf:
            if p.cpf in seen:
                p.errors.append(f"Duplicated CPF (row {seen[p.cpf]})")
            else:
                seen[p.cpf] = number
    return people


# ---------------------------------------------------------------- file generation

def _csv_row(p, password):
    return [p.cpf, p.firstname, p.lastname, p.email, password,
            f"{p.org}/{p.unit.upper()}", p.course, p.unit.upper()]


def generate_csv(people, password, delimiter=";"):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\r\n")
    writer.writerow(CSV_HEADER)
    for p in people:
        writer.writerow(_csv_row(p, password))
    return buffer.getvalue().encode("utf-8")


def generate_pdf(people, password, title):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=title,
                            leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm)
    styles = getSampleStyleSheet()
    cell = styles["BodyText"]
    cell.fontSize = 9
    cell.leading = 11

    data = [["UNIT", "NAME", "E-MAIL", "LOGIN", "PASSWORD"]]
    for p in people:
        data.append([Paragraph(p.unit, cell), Paragraph(p.name, cell),
                     Paragraph(p.email, cell), p.cpf, password])

    table = Table(data, colWidths=[6 * cm, 7.5 * cm, 7.5 * cm, 3.3 * cm, 2.6 * cm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f4c81")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa7b4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f4f8")]),
    ]))
    doc.build([
        Paragraph(title, styles["Title"]),
        Paragraph(f"Organization: {people[0].org} · Course: {people[0].course} · "
                  f"Generated on {date.today():%Y-%m-%d} · {len(people)} user(s)",
                  styles["Normal"]),
        Spacer(1, 0.4 * cm),
        table,
    ])
    return buffer.getvalue()


def _fill_sheet(ws, header, rows, widths):
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F4C81")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def generate_xlsx(people, password):
    wb = Workbook()
    ws = wb.active
    ws.title = "Step 1 - Input"
    _fill_sheet(ws, ["ORG CODE", "UNIT", "FULL NAME", "CPF", "E-MAIL", "COURSE"],
                [[p.org, p.unit, p.name, p.cpf, p.email, p.course] for p in people],
                [20, 35, 40, 16, 38, 30])
    for cell in ws["D"][1:]:
        cell.number_format = "@"  # CPF as text keeps the leading zeros

    _fill_sheet(wb.create_sheet("Step 2 - PDF list"), ["UNIT", "NAME", "E-MAIL", "LOGIN", "PASSWORD"],
                [[p.unit, p.name, p.email, p.cpf, password] for p in people],
                [35, 40, 38, 16, 12])
    _fill_sheet(wb.create_sheet("Step 3 - CSV"), CSV_HEADER,
                [_csv_row(p, password) for p in people],
                [16, 18, 30, 38, 12, 40, 30, 35])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- saving

def base_name(people):
    """COURSE_ORG_dd.mm.yyyy, based on the first row."""
    first = people[0]
    name = f"{first.course}_{first.org}_{date.today():%d.%m.%Y}"
    name = re.sub(r'[\\/:*?"<>|]', "-", name)
    return re.sub(r"-+", "-", name).strip()


def resolve_folder(template):
    """Replace {year} with the current year, e.g. output/pdf/{year} -> output/pdf/2026."""
    return Path(template.strip().replace("{year}", str(date.today().year)))


def free_name(folders_extensions, base):
    """Never overwrite existing files: append (2), (3)... when needed."""
    candidate, n = base, 1
    while any((folder / f"{candidate}{ext}").exists() for folder, ext in folders_extensions):
        n += 1
        candidate = f"{base} ({n})"
    return candidate


def save_files(people, folders, password, delimiter=";"):
    """folders: {"pdf": Path, "csv": Path, "xlsx": Path}. Returns {type: saved path}."""
    for folder in folders.values():
        folder.mkdir(parents=True, exist_ok=True)
    base = free_name([(f, f".{t}") for t, f in folders.items()], base_name(people))
    contents = {
        "csv": generate_csv(people, password, delimiter),
        "pdf": generate_pdf(people, password, "Moodle User Import"),
        "xlsx": generate_xlsx(people, password),
    }
    saved = {}
    for kind, content in contents.items():
        path = folders[kind] / f"{base}.{kind}"
        path.write_bytes(content)
        saved[kind] = path
    return saved


# ---------------------------------------------------------------- reading an uploaded spreadsheet

# header keywords (English and Portuguese) used to locate each column
SHEET_COLUMNS = {
    "org": ("ORG", "SIGLA"),
    "unit": ("UNIT", "UNIDADE"),
    "name": ("FULL NAME", "NOME COMPLETO"),
    "cpf": ("CPF",),
    "email": ("EMAIL", "E-MAIL"),
    "course": ("COURSE", "CURSO"),
}


def read_registration_sheet(file):
    """Read an .xlsx finding the header row by its text, wherever it is in the first 30 rows."""
    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    for idx, row in enumerate(rows[:30]):
        texts = [str(v).upper() if v is not None else "" for v in row]
        mapping = {}
        for field_name, keys in SHEET_COLUMNS.items():
            for col, text in enumerate(texts):
                if col not in mapping.values() and any(k in text for k in keys):
                    mapping[field_name] = col
                    break
        if {"name", "cpf", "email"} <= mapping.keys():
            break
    else:
        raise ValueError("Header row (FULL NAME, CPF, EMAIL) not found in the spreadsheet.")

    records = []
    for row in rows[idx + 1:]:
        record = {}
        for field_name, col in mapping.items():
            value = row[col] if col < len(row) else None
            record[field_name] = normalize_cpf(value) if field_name == "cpf" else ("" if value is None else str(value).strip())
        if any(record.values()):
            records.append(record)
    return records
